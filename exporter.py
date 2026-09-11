import io
from docx import Document
from fpdf import FPDF

class PDF(FPDF):
    def header(self):
        self.set_font('Arial', 'B', 10)
        self.set_text_color(100, 100, 100)
        self.cell(0, 10, 'Ata de Reunião - Gerado por IA Open Source', 0, 1, 'R')
        self.ln(5)

    def footer(self):
        self.set_y(-15)
        self.set_font('Arial', 'I', 8)
        self.set_text_color(150, 150, 150)
        self.cell(0, 10, f'Página {self.page_no()}', 0, 0, 'C')

def create_pdf_bytes(title, summary, transcript):
    """Gera um arquivo PDF em memória e retorna seus bytes."""
    pdf = PDF()
    pdf.add_page()
    
    pdf.set_font('Arial', 'B', 16)
    pdf.cell(0, 10, title, 0, 1, 'L')
    pdf.ln(5)
    
    pdf.set_font('Arial', 'B', 12)
    pdf.cell(0, 8, 'Ata e Resumo Executivo', 0, 1, 'L')
    pdf.set_font('Arial', '', 10)
    clean_summary = summary.encode('latin-1', 'replace').decode('latin-1')
    pdf.multi_cell(0, 6, clean_summary)
    pdf.ln(10)
    
    pdf.set_font('Arial', 'B', 12)
    pdf.cell(0, 8, 'Transcrição Completa', 0, 1, 'L')
    pdf.set_font('Arial', '', 9)
    clean_transcript = transcript.encode('latin-1', 'replace').decode('latin-1')
    pdf.multi_cell(0, 5, clean_transcript)
    
    return pdf.output(dest='S').encode('latin-1')

def create_docx_bytes(title, summary, transcript):
    """Gera um arquivo Word (.docx) em memória e retorna seus bytes."""
    doc = Document()
    
    doc.add_heading(title, level=1)
    
    doc.add_heading('Ata e Resumo Executivo', level=2)
    doc.add_paragraph(summary)
    
    doc.add_page_break()
    
    doc.add_heading('Transcrição Completa', level=2)
    doc.add_paragraph(transcript)
    
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()